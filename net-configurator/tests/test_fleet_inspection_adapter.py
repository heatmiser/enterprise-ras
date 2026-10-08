# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Focused fleet preparation and entry-boundary checks; no physical services."""

import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess

import pytest
import yaml
from tests.test_physical_endpoint_naming import MAPPING, example_inputs
import endpoint_naming

ROOT = Path(__file__).resolve().parent.parent
NAMES = ["server-control-01", "server-gpu-01"]
RUN = "20000101T120000Z"
spec = importlib.util.spec_from_file_location("fleet_adapter", ROOT / "scripts/prepare-fleet-inspection.py")
adapter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(adapter)


@pytest.fixture
def workspace(tmp_path):
    """Build an isolated example workbook/inventory; no site artifacts or credentials."""
    import openpyxl
    root = tmp_path / "net-configurator"
    settings, site, _ = example_inputs()
    settings.pop("node_endpoints")
    input_dir = root / "input/2-8-5-200/test01"
    inventory = root / "output/2-8-5-200/test01/inventory"
    input_dir.mkdir(parents=True)
    (inventory / "group_vars/all").mkdir(parents=True)
    (inventory / "host_vars").mkdir()
    (input_dir / "ocp-settings.yml").write_text(yaml.safe_dump(settings))
    (inventory / "group_vars/all/main.yml").write_text(yaml.safe_dump(site))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Settings"
    sheet.append(["OPENSHIFT"])
    sheet.append(["Setting", "Value"])
    sheet.append(["ocp_version", settings["cluster"]["version"]])
    nodes = workbook.create_sheet("Nodes")
    nodes.append(["Function", "Name", "OOB VLAN", "Type", "MAC Address for ZTP",
                  "Mgmt IP Address", "Prefix", "Gateway", "ZTP", "Enabled",
                  "Notes", "Include in Initial ABI"])
    wires = workbook.create_sheet("Wire Map")
    wires.append(["Display in Air", "System Name (A)", "Port (A)", "Port Side (A)",
                  "Cable Split (A)", "System Name (B)", "Port (B)", "Port Side (B)",
                  "Cable Split (B)", "Network Profile", "Kernel NIC Alias", "MAC Address"])
    for index, (name, device) in enumerate(site["devices"].items(), 1):
        nodes.append(["gpu" if "-gpu-" in name else "support", name, 200, "server",
                      device["mac"], device["eth0_ip"], 24, "192.0.2.1", "No", "Yes",
                      None, device["include_in_initial_abi"]])
        (inventory / "host_vars" / f"{name}.yml").write_text(
            yaml.safe_dump({"ansible_host": device["eth0_ip"]}))
        for profile, entries in device["nic_alias_map"].items():
            for nic, entry in enumerate(entries):
                switch = f"oob-{nic + 1:02}" if profile == "oob" else f"core-{nic % 2 + 1:02}"
                port = f"swp{index}s{nic // 2}" if profile != "oob" else f"swp{index}"
                if profile == "gpu": port = f"swp{index + 10}s{nic // 2}"
                wires.append(["Yes" if profile == "cpu" else "No", name, entry["alias"],
                              None, None, switch, port, None, None, profile,
                              entry["alias"], entry["mac"]])
        wires.append(["No", name, "iDRAC", None, None, "oob-01", f"swp{index + 20}",
                      None, None, "oob", "bmc", device["mac"]])
    vlans = workbook.create_sheet("VLANs & Profiles")
    vlans.append(["VLANs"])
    vlans.append(["VLAN ID", "Name", "Purpose", "Subnet", "Gateway", "VRF"])
    vlans.append([300, "cpu_network", "CPU", site["common"]["cpu_network"],
                  site["common"]["cpu_gateway"], "INBAND"])
    workbook.save(input_dir / "2-8-5-200.xlsx")
    workbook.close()
    collection = tmp_path / "collection"
    config = collection / "inventories/test01/cluster-vars.yaml"
    config.parent.mkdir(parents=True)
    config.write_text(yaml.safe_dump({"ocp_ee_image": "fixture-ee",
        "ocp_nodes": [{"name": name, "bmc": {"driver": "idrac", "host": site["devices"][name]["eth0_ip"]}}
                      for name in NAMES]}))
    return root, collection


def prepare(workspace):
    root, collection = workspace
    return adapter.prepare(root, "2-8-5-200", "test01", NAMES, RUN, collection, 2)


def seal(bundle):
    document = adapter.read_yaml(bundle / "plan.yml")
    node = document["inspect_fleet_plan"]["authorized_nodes"][0]
    adapter.write_yaml(bundle / "preflight-vars.yaml", {
        "inspection_preflight_provenance": {
            "installer_version": document["fleet_adapter_expected_version"],
            "callback_target_ip": node["devices"]["bond_ip"]},
        "preflight_inspection_nodes": [node],
        "ironic_image_resolve_ocp_release_image": "fixture-release",
        "ironic_image_resolve_oc_container_image": "fixture-ee",
        "ironic_deploy_provisioning_ip": "192.0.2.10",
        "ironic_deploy_provisioning_interface": "fixture0",
        "ironic_preprovisioning_media_nmstate_dir": str(bundle / "nmstate"),
        "ironic_preprovisioning_media_probe_node": node["name"],
        "ironic_preprovisioning_media_early_network_path": str(
            bundle / "early-network" / f"{node['name']}.yaml"),
        "ironic_preprovisioning_media_artifacts": {"fixture": "no services"}})
    adapter.write_yaml(bundle / "callback-routes.yaml", {"fixture": "No real route queries executed"})
    adapter.seal(bundle)


def test_example_cpu_identities_and_desired_ports_survive_preparation(workspace):
    bundle = prepare(workspace)
    document = adapter.read_yaml(bundle / "plan.yml")
    plan = document["inspect_fleet_plan"]
    assert [node["name"] for node in plan["authorized_nodes"]] == NAMES
    assert plan["expected_node_count"] == plan["concurrency_limit"] == 2
    assert len(plan["wiremap"]) == 12
    assert {purpose: sum(link["purpose"] == purpose for link in plan["wiremap"])
            for purpose in ("cpu", "gpu", "host_oob")} == {"cpu": 4, "gpu": 4, "host_oob": 4}
    k8s, gpu = plan["authorized_nodes"]
    assert k8s["rootDeviceHints"]["deviceName"] == "/dev/disk/by-path/pci-0000:01:00.0-nvme-1"
    assert gpu["rootDeviceHints"]["deviceName"] == "/dev/disk/by-path/pci-0000:81:00.0-nvme-1"
    assert [(link["interface"], link["expected_switch"], link["expected_port"])
            for link in plan["wiremap"] if link["node"] == NAMES[0] and link["purpose"] == "cpu"] == [
        ("ns-nic0", "core-01", "swp1s0"),
        ("ns-nic1", "core-02", "swp1s0")]
    assert {link["interface"] for link in plan["wiremap"] if link["purpose"] == "host_oob"} == {
        "host-oob0", "host-oob1"}
    assert {link["interface"] for link in plan["wiremap"] if link["purpose"] == "gpu"} == {
        "rail0", "rail1", "rail2", "rail3"}
    assert all(link["mac"] not in {"02:ff:00:00:00:01", "02:ff:00:00:00:04"}
               for link in plan["wiremap"])  # BMC endpoints are not IPA host NICs.
    assert "storagePolicy" not in gpu  # Never inherit K8S controller/volume IDs.
    assert not Path(plan["artifact_root"]).exists()
    assert not (workspace[0] / "output/2-8-5-200/test01/ocp/agent-config.yaml").exists()
    for node in plan["authorized_nodes"]:
        nmstate = adapter.read_yaml(bundle / "nmstate" / f"{node['name']}.yaml")
        early = adapter.read_yaml(bundle / "early-network" / f"{node['name']}.yaml")
        assert [interface["name"] for interface in nmstate["interfaces"]] == ["bond0"]
        assert nmstate["interfaces"][0]["link-aggregation"]["port"] == ["ns-nic0", "ns-nic1"]
        assert early["interfaces"] == [{"name": interface["name"], "mac": interface["macAddress"]}
                                       for interface in node["interfaces"]]


def test_preparation_refuses_stale_mac_before_output(workspace):
    path = workspace[0] / "output/2-8-5-200/test01/inventory/group_vars/all/main.yml"
    document = adapter.read_yaml(path)
    document["devices"][NAMES[0]]["nic_alias_map"]["cpu"][0]["mac"] = "00:11:22:33:44:55"
    path.write_text(yaml.safe_dump(document))
    with pytest.raises(ValueError, match="Workbook CPU identity differs"):
        prepare(workspace)
    assert not (workspace[0] / "output/2-8-5-200/test01/ocp/inspection/fleet" / RUN).exists()


@pytest.mark.parametrize("profile", ["gpu", "oob"])
def test_preparation_refuses_stale_auxiliary_mac_before_output(workspace, profile):
    path = workspace[0] / "output/2-8-5-200/test01/inventory/group_vars/all/main.yml"
    document = adapter.read_yaml(path)
    document["devices"][NAMES[1]]["nic_alias_map"][profile][0]["mac"] = "02:00:00:00:ee:ee"
    path.write_text(yaml.safe_dump(document))
    with pytest.raises(ValueError, match=f"Workbook {profile} identity differs"):
        prepare(workspace)
    assert not (workspace[0] / "output/2-8-5-200/test01/ocp/inspection/fleet" / RUN).exists()


def test_revised_workbook_gpu_links_match_recorded_physical_ports():
    """Optional operator-supplied physical evidence; identifiers stay outside the repository."""
    configured = [os.environ.get(key) for key in (
        "FLEET_TEST_WORKBOOK", "FLEET_TEST_REPORT", "FLEET_TEST_NODE")]
    if not all(configured):
        pytest.skip("Supply FLEET_TEST_WORKBOOK, FLEET_TEST_REPORT and FLEET_TEST_NODE")
    workbook, report_path = map(Path, configured[:2])
    node_name = configured[2]
    assert workbook.is_file() and report_path.is_file()
    wb = adapter.excel_parser.load_workbook_safe(workbook, data_only=True)
    try:
        wires = adapter.excel_parser._build_wiremap_row_list(wb["Wire Map"])
    finally:
        wb.close()
    device = {"nic_alias_map": adapter.excel_parser.build_nic_alias_map(wires, node_name)}
    links = adapter.auxiliary_links(node_name, wires, device)
    assert links  # Validate all supplied auxiliary links without assuming a rail count.
    raw = yaml.safe_load(report_path.read_text())["raw_ironic_responses"][node_name]
    ports = raw["ports"]["ports"]
    nics = raw["inventory"]["inventory"]["interfaces"]
    for link in links:
        matching_ports = [port for port in ports if port["address"] == link["mac"]]
        matching_nics = [nic for nic in nics if nic["mac_address"] == link["mac"] and nic["pci_address"]]
        assert len(matching_ports) == len(matching_nics) == 1
        assert matching_nics[0]["has_carrier"] is True
        assert matching_ports[0]["local_link_connection"]["switch_info"] == link["expected_switch"]
        assert matching_ports[0]["local_link_connection"]["port_id"] == link["expected_port"]


def test_preparation_refuses_missing_bmc_declaration(workspace):
    path = workspace[1] / "inventories/test01/cluster-vars.yaml"
    document = adapter.read_yaml(path)
    document["ocp_nodes"].pop()
    path.write_text(yaml.safe_dump(document))
    with pytest.raises(ValueError, match="unique workbook/collection node"):
        prepare(workspace)


@pytest.mark.parametrize("matching", [True, False])
def test_named_bmc_endpoints_keep_physical_authority_and_verify_resolution(workspace, monkeypatch, matching):
    root, collection = workspace
    settings_path = root / "input/2-8-5-200/test01/ocp-settings.yml"
    settings = adapter.read_yaml(settings_path)
    settings["cluster"].update(name="cluster-01", domain="example.com")
    settings["node_endpoints"] = MAPPING
    settings_path.write_text(yaml.safe_dump(settings))
    cluster_path = collection / "inventories/test01/cluster-vars.yaml"
    cluster = adapter.read_yaml(cluster_path)
    for node in cluster["ocp_nodes"]:
        node["bmc"]["host"] = MAPPING[node["name"]]["bmc_fqdn"]
    cluster_path.write_text(yaml.safe_dump(cluster))
    expected = {MAPPING[NAMES[0]]["bmc_fqdn"]: "192.0.2.101",
                MAPPING[NAMES[1]]["bmc_fqdn"]: "192.0.2.201"}
    looked_up = []

    def resolve(host, *args, **kwargs):
        looked_up.append(host)
        ip = expected[host] if matching else "192.0.2.99"
        return [(2, 1, 6, "", (ip, 0))]

    monkeypatch.setattr(endpoint_naming.socket, "getaddrinfo", resolve)
    if not matching:
        with pytest.raises(ValueError, match="BMC DNS does not match workbook IP"):
            prepare(workspace)
        assert not (root / "output/2-8-5-200/test01/ocp/inspection/fleet" / RUN).exists()
        return
    bundle = prepare(workspace)
    plan = adapter.read_yaml(bundle / "plan.yml")["inspect_fleet_plan"]
    assert [node["name"] for node in plan["authorized_nodes"]] == NAMES
    assert [node["bmc"]["host"] for node in plan["authorized_nodes"]] == list(expected)
    assert looked_up == list(expected)
    assert [node["devices"]["bond_ip"] for node in plan["authorized_nodes"]] == ["198.51.100.101", "198.51.100.201"]


@pytest.mark.parametrize("case", ["matching", "wrong_run", "wrong_nodes", "changed_network", "existing_evidence"])
def test_exact_authority_and_sealed_inputs(workspace, case):
    bundle = prepare(workspace)
    seal(bundle)
    run, names = RUN, NAMES
    if case == "wrong_run":
        run = "20000101T120001Z"
    elif case == "wrong_nodes":
        names = NAMES[:1]
    elif case == "changed_network":
        (bundle / "nmstate" / f"{NAMES[0]}.yaml").write_text("changed input")
    elif case == "existing_evidence":
        Path(adapter.read_yaml(bundle / "plan.yml")["inspect_fleet_plan"]["artifact_root"]).mkdir(parents=True)
    if case == "matching":
        auth = adapter.authorize(bundle, run, names)["fleet_inspection_driver_physical_authorization"]
        assert auth == {"confirmed": True, "action": "inspect", "run_id": RUN, "node_names": NAMES}
    else:
        with pytest.raises(ValueError):
            adapter.authorize(bundle, run, names)


def test_make_handoff_and_authorization_guard_without_ansible_services(workspace):
    root, collection = workspace
    shutil.copyfile(ROOT / "Makefile", root / "Makefile")
    shutil.copytree(ROOT / "scripts", root / "scripts")
    bundle = prepare(workspace)
    seal(bundle)
    (collection / "inventories/test01/secrets.yaml").write_text("unused fixture")
    bin_dir = root / "bin"
    bin_dir.mkdir()
    capture = root / "captured.json"
    executable = bin_dir / "ansible-playbook"
    executable.write_text("#!/usr/bin/env python3\nimport json, os, sys\nfrom pathlib import Path\n"
                          "Path(os.environ['FLEET_TEST_CAPTURE']).write_text(json.dumps(sys.argv[1:]))\n")
    executable.chmod(0o700)
    command = ["make", "inspect-fleet", "ARCH=2-8-5-200", "SITE=test01", f"FLEET_RUN_ID={RUN}",
               f"FLEET_AUTHORIZE_RUN={RUN}", f"RHV_BAREMETAL_OCP_ROOT={collection}"]
    env = {**os.environ, "PATH": f"{bin_dir}:{os.environ['PATH']}", "FLEET_TEST_CAPTURE": str(capture)}
    rejected = subprocess.run(command + [f"FLEET_AUTHORIZE_NODES={NAMES[0]}"], cwd=root,
                              env=env, capture_output=True, text=True, timeout=60)
    assert rejected.returncode != 0 and not capture.exists()
    accepted = subprocess.run(command + [f"FLEET_AUTHORIZE_NODES={','.join(NAMES)}"], cwd=root,
                              env=env, capture_output=True, text=True, timeout=60)
    assert accepted.returncode == 0, accepted.stdout + accepted.stderr
    arguments = json.loads(capture.read_text())
    assert "playbooks/inspect-fleet.yml" in arguments
    assert f"@{bundle}/hosts.yml" not in arguments
    assert str(bundle / "hosts.yml") in arguments
    assert f"@{bundle}/preflight-vars.yaml" not in arguments
    authority = next(json.loads(arg) for arg in arguments if arg.startswith('{"fleet_inspection_driver_physical_authorization"'))
    assert authority["fleet_inspection_driver_physical_authorization"]["node_names"] == NAMES


@pytest.mark.parametrize("authorized", [True, False])
def test_actual_adapter_guard_with_harmless_collection_entry(workspace, authorized):
    """Run production per-node network preparation; stop before runtime or BMC operations."""
    root, _ = workspace
    bundle = prepare(workspace)
    seal(bundle)
    shutil.copytree(ROOT / "scripts", root / "scripts")
    playbooks = root / "playbooks"
    playbooks.mkdir()
    marker = root / "entered.yaml"
    stub = playbooks / "collection-fixture.yml"
    stub.write_text(yaml.safe_dump([{
        "name": "Prepare selected networks without any service or hardware",
        "hosts": "bootstrap", "gather_facts": False,
        "tasks": [{"name": "Validate and prepare real per-node collection networking",
                   "ansible.builtin.include_role": {
                       "name": "rhvp.baremetal_ocp.fleet_inspection_driver", "tasks_from": "validate_plan"},
                   "vars": {"fleet_inspection_driver_prepare_network": True}},
                  {"name": "Record exact selected handoff", "ansible.builtin.copy": {
            "dest": str(marker), "mode": "0600",
            "content": "{{ {'names': fleet_inspection_driver_bmc_credentials.keys() | list, "
                       "'networks': __fleet_inspection_driver_prepared_nodes, "
                       "'authority': fleet_inspection_driver_physical_authorization} | to_nice_yaml }}"}}]}]))
    play = playbooks / "inspect-fleet.yml"
    play.write_text((ROOT / "playbooks/inspect-fleet.yml").read_text().replace(
        "ansible.builtin.import_playbook: rhvp.baremetal_ocp.inspect_fleet",
        "ansible.builtin.import_playbook: collection-fixture.yml"))
    inputs = {"fleet_adapter_authorized_run": RUN,
              "fleet_adapter_authorized_nodes": NAMES if authorized else NAMES[:1],
              "bmc_credentials": {name: {"username": "fixture-user", "password": "fixture-password"}
                                  for name in NAMES + ["unselected-node"]}}
    inputs_path = root / "inputs.json"
    inputs_path.write_text(json.dumps(inputs))
    result = subprocess.run(["ansible-playbook", "-i", str(bundle / "hosts.yml"), str(play),
                             "-e", f"@{bundle}/plan.yml",
                             "-e", f"@{inputs_path}"], text=True, capture_output=True, timeout=60)
    if authorized:
        assert result.returncode == 0, result.stdout + result.stderr
        handoff = yaml.safe_load(marker.read_text())
        assert handoff["names"] == NAMES
        assert handoff["authority"]["node_names"] == NAMES
        nodes = adapter.read_yaml(bundle / "plan.yml")["inspect_fleet_plan"]["authorized_nodes"]
        networks = handoff["networks"]
        assert set(networks) == set(NAMES)
        assert [networks[name]["callback_ip"] for name in NAMES] == ["198.51.100.101", "198.51.100.201"]
        for node in nodes:
            network = networks[node["name"]]
            assert network["nmstate_path"] == str(bundle / "nmstate" / f"{node['name']}.yaml")
            assert network["early_network_path"] == str(bundle / "early-network" / f"{node['name']}.yaml")
            assert f"ip={node['devices']['bond_ip']}::" in network["kernel_append_params"]
            for interface in node["interfaces"]:
                assert f"ifname={interface['name']}:{interface['macAddress']}" in network["kernel_append_params"]
            other = next(candidate for candidate in nodes if candidate["name"] != node["name"])
            assert all(interface["macAddress"] not in network["kernel_append_params"]
                       for interface in other["interfaces"])
        assert not Path(adapter.read_yaml(bundle / "plan.yml")["inspect_fleet_plan"]["artifact_root"]).exists()
    else:
        assert result.returncode != 0
        assert not marker.exists()
