"""Initial ABI membership and Day-2 node image render contract."""

import importlib.util
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "generate-ocp-inventory.py"
sys.path.insert(0, str(SCRIPT.parent))
from excel_parser import build_devices, parse_nodes
from validate_excel import ValidationResult, validate_nodes


spec = importlib.util.spec_from_file_location("generate_ocp_inventory_membership", SCRIPT)
gen_ocp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(gen_ocp)


def test_workbook_membership_is_header_driven_and_control_plane_is_required():
    wb = Workbook()
    ws = wb.active
    ws.title = "Nodes"
    ws.append(["Function", "Name", "Mgmt IP Address", "Include in Initial ABI"])
    ws.append(["support", "cp-01", "192.0.2.11", "Yes"])
    ws.append(["gpu", "gpu-03", "192.0.2.13", "No"])

    nodes = parse_nodes(ws)
    assert [n["include_in_initial_abi"] for n in nodes] == ["Yes", "No"]
    devices = build_devices(nodes, [], [])
    assert devices["gpu-03"]["include_in_initial_abi"] == "No"

    result = ValidationResult()
    validate_nodes(ws, result, ocp_hostnames={"cp-01": "control_plane", "gpu-03": "worker_gpu"})
    assert not any("Include in Initial ABI" in str(error) for error in result.errors)

    ws["D2"] = "No"
    result = ValidationResult()
    validate_nodes(ws, result, ocp_hostnames={"cp-01": "control_plane", "gpu-03": "worker_gpu"})
    assert any("control-plane node" in str(error) for error in result.errors)

    ws["D2"] = "Yes"
    ws["D3"] = None
    result = ValidationResult()
    validate_nodes(ws, result, ocp_hostnames={"cp-01": "control_plane", "gpu-03": "worker_gpu"})
    assert any("must be exactly Yes or No" in str(error) for error in result.errors)

    ws.delete_cols(4)
    result = ValidationResult()
    validate_nodes(ws, result, ocp_hostnames={"cp-01": "control_plane"})
    assert any("Missing required column header" in str(error) for error in result.errors)


@pytest.mark.parametrize("value", [None, "", "yes", "Maybe"])
def test_membership_rejects_missing_or_noncanonical_values(value):
    with pytest.raises(ValueError, match="must be exactly Yes or No"):
        gen_ocp.split_abi_membership(
            {"gpu-03": "worker_gpu"},
            {"gpu-03": {"include_in_initial_abi": value}},
        )


def test_day2_worker_reuses_identity_disk_and_north_south_network(tmp_path, monkeypatch):
    disk = "/dev/disk/by-path/pci-0000:81:00.0-nvme-1"
    worker = {
        "include_in_initial_abi": "No",
        "mac": "00:11:22:33:44:00",
        "bond_ip": "10.78.221.203/24",
        "nic_alias_map": {
            "cpu": [
                {"alias": "ns-nic0", "mac": "00:11:22:33:44:01"},
                {"alias": "ns-nic1", "mac": "00:11:22:33:44:02"},
            ],
            "gpu": [{"alias": "rail0", "mac": "00:11:22:33:44:03"}],
        },
    }
    control = dict(worker, include_in_initial_abi="Yes", bond_ip="10.78.221.101/24")
    devices = {"cp-01": control, "gpu-03": worker}
    roles = {"cp-01": "control_plane", "gpu-03": "worker_gpu"}
    initial, day2 = gen_ocp.split_abi_membership(roles, devices)
    assert initial == {"cp-01": "control_plane"}
    assert day2 == {"gpu-03": "worker_gpu"}

    site = {"devices": devices, "common": {"cpu_gateway": "10.78.221.1", "cpu_network": "10.78.221.0/24"},
            "dns_servers": ["192.0.2.53"]}
    settings = {"cluster": {"name": "test", "domain": "example.com"}, "oem": "dell"}
    agent = gen_ocp.build_agent_config(settings, initial, {}, site, "2-8-5-200")
    assert [host["hostname"] for host in agent["hosts"]] == ["cp-01"]
    assert agent["rendezvousIP"] == "10.78.221.101"

    rendered = gen_ocp.build_day2_workers_config(settings, day2, {}, site, "2-8-5-200")
    assert list(rendered) == ["hosts"]
    host = rendered["hosts"][0]
    assert host["hostname"] == "gpu-03"
    assert "role" not in host
    assert host["rootDeviceHints"] == {"deviceName": disk}
    assert {item["name"]: item["macAddress"] for item in host["interfaces"]}["ns-nic0"] == "00:11:22:33:44:01"
    assert [item["name"] for item in host["networkConfig"]["interfaces"]] == ["ns-nic0", "ns-nic1", "ns-bond0"]

    key = tmp_path / "key.pub"
    key.write_text("ssh-ed25519 synthetic-test-key")
    settings["ssh_key_path"] = str(key)
    monkeypatch.setattr(gen_ocp, "read_pull_secret", lambda _settings, _vault: ("synthetic", "test"))
    install = gen_ocp.build_install_config(settings, site, initial)
    assert install["compute"][0]["replicas"] == 0
    assert install["controlPlane"]["replicas"] == 1


@pytest.mark.parametrize("missing_credential", ["pull_secret", "ssh_key"])
def test_generation_preserves_install_config_when_credentials_unavailable(
    tmp_path, monkeypatch, capsys, missing_credential
):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", [str(SCRIPT), "--arch", "2-8-5-200", "--site", "test01"])
    monkeypatch.setattr(gen_ocp, "_load_shared_air_vault", lambda: {})
    monkeypatch.setattr(gen_ocp, "load_disk_defaults", lambda: {})

    roles = {
        "control_plane": [f"cp-{number:02d}" for number in range(1, 4)],
        "worker_gpu": [f"gpu-{number:02d}" for number in range(1, 5)],
    }
    devices = {}
    for index, hostname in enumerate(roles["control_plane"] + roles["worker_gpu"], start=1):
        devices[hostname] = {
            "include_in_initial_abi": "No" if hostname == "gpu-03" else "Yes",
            "mac": f"02:00:00:00:00:{index:02x}",
            "bond_ip": f"192.0.2.{index}/24",
        }
    site = {"devices": devices, "common": {
        "cpu_gateway": "192.0.2.254", "cpu_network": "192.0.2.0/24",
    }}
    inventory = tmp_path / "output/2-8-5-200/test01/inventory/group_vars/all/main.yml"
    inventory.parent.mkdir(parents=True)
    inventory.write_text(yaml.safe_dump(site))

    pull_secret = tmp_path / "pull-secret.json"
    ssh_key = tmp_path / "key.pub"
    if missing_credential != "pull_secret":
        pull_secret.write_text('{"auths": {}}')
    if missing_credential != "ssh_key":
        ssh_key.write_text("ssh-ed25519 synthetic-test-key")
    settings = {
        "cluster": {"name": "test", "domain": "example.com"},
        "node_roles": roles,
        "pull_secret_path": str(pull_secret),
        "ssh_key_path": str(ssh_key),
    }
    settings_path = tmp_path / "input/2-8-5-200/test01/ocp-settings.yml"
    settings_path.parent.mkdir(parents=True)
    settings_path.write_text(yaml.safe_dump(settings))

    output = tmp_path / "output/2-8-5-200/test01/ocp"
    output.mkdir()
    install_path = output / "install-config.yaml"
    original = b"# Existing deployment manifest\nmetadata:\n  name: preserved-cluster\n"
    install_path.write_bytes(original)

    gen_ocp.main()

    assert install_path.read_bytes() == original
    assert "Preserved existing" in capsys.readouterr().out
    agent = yaml.safe_load((output / "agent-config.yaml").read_text())
    assert {host["hostname"] for host in agent["hosts"]} == set(devices) - {"gpu-03"}
    day2 = yaml.safe_load((output / "day2/workers/nodes-config.yaml").read_text())
    assert [host["hostname"] for host in day2["hosts"]] == ["gpu-03"]
    generated_inventory = yaml.safe_load((output / "inventory/hosts.yml").read_text())
    groups = generated_inventory["all"]["children"]
    assert set(groups["ocp_worker_gpu"]["hosts"]) == set(roles["worker_gpu"])
    assert set(groups["ocp_control_plane"]["hosts"]) == set(roles["control_plane"])
