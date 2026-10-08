# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Physical endpoint contracts over self-contained example device/interface inputs."""

import base64
import copy
import importlib.util
from pathlib import Path
import sys

import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from endpoint_naming import (build_endpoint_manifest, render_dnsmasq_records,
                             build_role_endpoints, WORKBOOK_ROLE_TO_OCP)


def module(filename):
    spec = importlib.util.spec_from_file_location(filename.replace("-", "_"), ROOT / "scripts" / filename)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


GEN = module("generate-ocp-inventory.py")
INIT = module("init-ocp-settings.py")
CLUSTER = "cluster-01.example.com"
MAPPING = {
    "server-control-01": {"bmc_fqdn": "control-01.mgmt-01.example.com", "rhcos_fqdn": "control-01." + CLUSTER},
    "server-control-02": {"bmc_fqdn": "control-02.mgmt-01.example.com", "rhcos_fqdn": "control-02." + CLUSTER},
    "server-control-03": {"bmc_fqdn": "control-03.mgmt-01.example.com", "rhcos_fqdn": "control-03." + CLUSTER},
    "server-gpu-01": {"bmc_fqdn": "gpu-01.mgmt-01.example.com", "rhcos_fqdn": "gpu-01." + CLUSTER},
    "server-gpu-02": {"bmc_fqdn": "gpu-02.mgmt-01.example.com", "rhcos_fqdn": "gpu-02." + CLUSTER},
    "server-gpu-03": {"bmc_fqdn": "gpu-03.mgmt-01.example.com", "rhcos_fqdn": "gpu-03." + CLUSTER},
    "server-gpu-04": {"bmc_fqdn": "gpu-04.mgmt-01.example.com", "rhcos_fqdn": "gpu-04." + CLUSTER},
}


def example_inputs():
    """Representative topology using documentation IPs and locally administered MACs."""
    controls = [f"server-control-{i:02}" for i in range(1, 4)]
    workers = [f"server-gpu-{i:02}" for i in range(1, 5)]
    settings = {
        "cluster": {"name": "cluster-01", "domain": "example.com", "version": "4.22.14",
                    "api_vip": "198.51.100.254", "ingress_vip": "198.51.100.253"},
        "oem": "dell", "install_disk": {"overrides": {}},
        "node_roles": {"control_plane": controls, "worker_gpu": workers},
        "node_endpoints": copy.deepcopy(MAPPING),
    }
    site = {"architecture": "2-8-5-200", "dns_servers": ["192.0.2.61"],
            "common": {"cpu_network": "198.51.100.0/24", "cpu_gateway": "198.51.100.1"},
            "devices": {}}
    for index, name in enumerate(controls + workers, 1):
        suffix = int(name.rsplit("-", 1)[1])
        address = 100 + suffix if name in controls else 200 + suffix
        aliases = {profile: [{"alias": alias, "mac": f"02:00:00:{index:02x}:{kind:02x}:{nic:02x}"}
                             for nic, alias in enumerate(names)]
                   for kind, (profile, names) in enumerate([
                       ("cpu", ["ns-nic0", "ns-nic1"]), ("oob", ["host-oob0", "host-oob1"]),
                       ("gpu", [f"rail{i}" for i in range(4)] if name in workers else [])], 1)}
        device = {"eth0_ip": f"192.0.2.{address}", "mac": f"02:ff:00:00:00:{index:02x}",
                  "bond_ip": f"198.51.100.{address}/24", "nic_alias_map": aliases,
                  "interfaces": {profile: [entry["alias"] for entry in entries]
                                 for profile, entries in aliases.items()},
                  "include_in_initial_abi": "No" if name == "server-gpu-03" else "Yes"}
        if name in workers:
            device["gpu_ips"] = [f"203.0.113.{10 * index + rail}/24" for rail in range(4)]
        site["devices"][name] = device
    roles = GEN.build_role_map(settings, site["devices"], "2-8-5-200", "example-site")
    return settings, site, roles


@pytest.fixture
def inputs():
    return example_inputs()


def test_reviewed_addresses_generate_node_and_cluster_dns_records(inputs):
    settings, site, roles = inputs
    manifest = build_endpoint_manifest(settings, site["devices"], roles)
    assert set(manifest["nodes"]) == set(MAPPING)
    assert manifest["nodes"]["server-control-01"] == {
        "bmc": {"hostname": MAPPING["server-control-01"]["bmc_fqdn"], "ip": "192.0.2.101"},
        "rhcos": {"hostname": MAPPING["server-control-01"]["rhcos_fqdn"], "ip": "198.51.100.101"},
    }
    text = render_dnsmasq_records(manifest)
    assert f"host-record=gpu-03.{CLUSTER},198.51.100.203\n" in text
    assert f"host-record=api.{CLUSTER},api-int.{CLUSTER},198.51.100.254\n" in text
    assert f"address=/apps.{CLUSTER}/198.51.100.253\n" in text
    assert "local=/mgmt-01.example.com/\n" in text
    assert manifest["rhcos_to_physical"]["control-01." + CLUSTER] == "server-control-01"


@pytest.mark.parametrize("case", ["missing_node", "duplicate_name", "wrong_cluster"])
def test_mapping_errors_fail_before_rendering(inputs, case):
    settings, site, roles = inputs
    if case == "missing_node":
        settings["node_endpoints"].pop("server-gpu-03")
    elif case == "duplicate_name":
        settings["node_endpoints"]["server-gpu-02"] = dict(MAPPING["server-gpu-01"])
    else:
        settings["cluster"]["name"] = "different-cluster"
    with pytest.raises(ValueError):
        build_endpoint_manifest(settings, site["devices"], roles)


@pytest.mark.parametrize("role", ["control_plane", "infra", "worker", "worker_gpu", "worker_storage"])
def test_conflicting_bond_fields_match_generated_network_address(inputs, role):
    settings, site, _ = inputs
    physical = "server-control-01"
    settings["node_endpoints"] = {physical: MAPPING[physical]}
    device = site["devices"][physical]
    device["bond_ip"] = "198.51.100.101/24"
    device["bond_ip1"] = "203.0.113.101/24"
    device["interfaces"]["storage"] = device["interfaces"]["cpu"]
    device["nic_alias_map"]["storage"] = device["nic_alias_map"]["cpu"]
    manifest = build_endpoint_manifest(settings, {physical: device}, {physical: role})
    config = GEN.build_nmstate_network_config(device, site, role)
    bond = next(interface for interface in config["interfaces"] if interface["type"] == "bond")
    address = bond["ipv4"]["address"][0]["ip"]
    assert manifest["nodes"][physical]["rhcos"]["ip"] == address
    assert f"host-record={MAPPING[physical]['rhcos_fqdn']},{address}\n" in render_dnsmasq_records(manifest)


def test_real_generation_separates_physical_keys_from_installed_names(inputs, tmp_path, monkeypatch):
    settings, site, roles = inputs
    inventory = tmp_path / "output/2-8-5-200/example-site/inventory"
    (inventory / "group_vars/all").mkdir(parents=True)
    (inventory / "host_vars").mkdir()
    for name, device in site["devices"].items():
        (inventory / "host_vars" / f"{name}.yml").write_text(
            yaml.safe_dump({"ansible_host": device["eth0_ip"]}))
    (inventory / "group_vars/all/main.yml").write_text(yaml.safe_dump(site))
    settings_file = tmp_path / "input/2-8-5-200/example-site/ocp-settings.yml"
    settings_file.parent.mkdir(parents=True)
    ssh_key = tmp_path / "test.pub"
    # Public-only RFC 8032 test vector; no credentials or private key material.
    public_bytes = bytes.fromhex("d75a980182b10ab7d54bfed3c964073a0ee172f3daa62325af021a68f707511a")
    key_blob = b"\x00\x00\x00\x0bssh-ed25519\x00\x00\x00\x20" + public_bytes
    ssh_key.write_text("ssh-ed25519 " + base64.b64encode(key_blob).decode() + " fixture\n")
    settings["ssh_key_path"] = str(ssh_key)
    settings_file.write_text(yaml.safe_dump(settings))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["generate-ocp-inventory.py", "--arch", "2-8-5-200", "--site", "example-site"])
    monkeypatch.setattr(GEN, "_load_shared_air_vault", lambda: {"ocp_pull_secret": '{"auths":{}}'})
    GEN.main()
    output = inventory.parent / "ocp"
    agent = GEN.load_yaml(output / "agent-config.yaml")
    assert {host["hostname"] for host in agent["hosts"]} == {
        MAPPING[name]["rhcos_fqdn"] for name in roles if name != "server-gpu-03"
    }
    assert all(host["networkConfig"]["dns-resolver"]["config"]["server"] == ["192.0.2.61"]
               for host in agent["hosts"])
    assert agent["rendezvousIP"] == "198.51.100.101"
    workers = GEN.load_yaml(output / "day2/workers/nodes-config.yaml")
    assert [host["hostname"] for host in workers["hosts"]] == [MAPPING["server-gpu-03"]["rhcos_fqdn"]]
    hosts = GEN.load_yaml(output / "inventory/hosts.yml")["all"]["children"]["ocp_control_plane"]["hosts"]
    assert set(hosts) == {name for name, role in roles.items() if role == "control_plane"}
    physical = "server-control-01"
    hv = GEN.load_yaml(output / f"inventory/host_vars/{physical}.yml")
    assert hv["physical_server_name"] == physical
    assert hv["ansible_host"] == MAPPING[physical]["rhcos_fqdn"]
    assert hv["bmc_ip"] == "192.0.2.101" and hv["rhcos_ip"] == "198.51.100.101"
    nncp = GEN.load_yaml(output / "day2/nncp-server-gpu-01-gpu-rails.yaml")
    assert nncp["spec"]["nodeSelector"]["kubernetes.io/hostname"] == MAPPING["server-gpu-01"]["rhcos_fqdn"]
    install = GEN.load_yaml(output / "install-config.yaml")
    assert install["baseDomain"] == "example.com" and install["metadata"]["name"] == "cluster-01"
    assert GEN.load_yaml(output / "endpoint-map.yaml")["nodes"][physical]["bmc"]["ip"] == hv["bmc_ip"]
    assert (output / "dns/dnsmasq-records.conf").is_file()


def test_initializer_preserves_mapping_and_explicit_cluster_name(inputs, tmp_path):
    settings, _, _ = inputs
    output = tmp_path / "ocp-settings.yml"
    output.write_text(yaml.safe_dump(settings))
    roles_yaml = yaml.safe_dump(settings["node_roles"], sort_keys=False)
    roles_yaml = "\n".join("  " + line for line in roles_yaml.splitlines())
    INIT.write_settings(output, "2-8-5-200", "example-site", roles_yaml,
                        ocp={"ocp_cluster_domain": "example.com"})
    regenerated = GEN.load_yaml(output)
    assert regenerated["node_endpoints"] == MAPPING
    assert regenerated["cluster"]["name"] == "cluster-01"
    before = output.read_bytes()
    with pytest.raises(ValueError, match="regenerated cluster domain"):
        INIT.write_settings(output, "2-8-5-200", "example-site", roles_yaml,
                            ocp={"ocp_cluster_name": "other", "ocp_cluster_domain": "example.com"})
    assert output.read_bytes() == before


def test_unconfigured_names_retain_existing_inventory_behavior(inputs):
    settings, site, roles = inputs
    settings.pop("node_endpoints")
    assert build_endpoint_manifest(settings, site["devices"], roles) is None
    agent = GEN.build_agent_config(settings, roles, {}, site, "2-8-5-200")
    assert {host["hostname"] for host in agent["hosts"]} == set(roles)
    inventory = GEN.build_hosts_yaml({"server-control-01": "control_plane"},
                                    {"server-control-01": {"ansible_host": "192.0.2.101"}})
    assert inventory["all"]["children"]["ocp_control_plane"]["hosts"]["server-control-01"] == {
        "ansible_host": "192.0.2.101"
    }


@pytest.mark.parametrize('role', list(WORKBOOK_ROLE_TO_OCP))
def test_role_names_preserve_suffix_and_role_label(role):
    entry = build_role_endpoints({'physical-server-007': role}, 'cluster-01', 'example.com', ' ')
    assert entry['physical-server-007'] == {
        'bmc_fqdn': f'{role}-007.mgmt-01.example.com',
        'rhcos_fqdn': f'{role}-007.cluster-01.example.com',
    }


@pytest.mark.parametrize('roles', [
    {'server01': 'control'},
    {'server-01': 'unsupported'},
    {'server-a-01': 'worker', 'server-b-01': 'worker'},
])
def test_role_endpoint_rejections(roles):
    with pytest.raises(ValueError):
        build_role_endpoints(roles, 'cluster-01', 'example.com')


def _role_workbook(path, assignments):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Nodes'
    ws.append(['Function', 'Name', 'Enabled', 'Include in Initial ABI', 'Role'])
    for function, name, role, enabled, membership in assignments:
        ws.append([function, name, enabled, membership, role])
    settings = wb.create_sheet('Settings')
    settings.append(['MANAGEMENT'])
    settings.append(['Setting', 'Value'])
    settings.append(['bmc_dns_subdomain', 'management'])
    settings.append([])
    settings.append(['OPENSHIFT'])
    settings.append(['Setting', 'Value'])
    settings.append(['ocp_cluster_name', 'cluster-01'])
    settings.append(['ocp_cluster_domain', 'example.com'])
    wb.save(path)
    wb.close()


def test_initializer_main_generates_all_roles_and_deferred_endpoints(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / 'input/demo/site/demo.xlsx'
    path.parent.mkdir(parents=True)
    _role_workbook(path, [
        ('support', 'server-cp-01', 'control', 'Yes', 'Yes'),
        ('gpu', 'server-gpu-02', 'gpu', 'Yes', 'No'),
        ('storage', 'server-storage-03', 'storage', 'Yes', 'Yes'),
        ('support', 'server-worker-04', 'worker', 'Yes', 'Yes'),
        ('support', 'server-infra-05', 'infra', 'Yes', 'Yes'),
        ('gpu', 'server-disabled-06', 'gpu', 'No', 'No'),
        ('utility', 'utility', None, 'Air', None),
    ])
    inventory = tmp_path / 'output/demo/site/inventory/hosts'
    inventory.parent.mkdir(parents=True)
    inventory.write_text('[support]\nserver-cp-01\nserver-worker-04\nserver-infra-05\n'
                         '[nodes]\nserver-gpu-02\n[storage]\nserver-storage-03\n')
    monkeypatch.setattr(INIT, '_load_shared_air_vault', lambda: {})
    monkeypatch.setattr(sys, 'argv', ['init-ocp-settings.py', '--arch', 'demo', '--site', 'site'])
    INIT.main()
    output = path.parent / 'ocp-settings.yml'
    result = yaml.safe_load(output.read_text())
    assert result['node_roles'] == {
        'control_plane': ['server-cp-01'], 'worker_gpu': ['server-gpu-02'],
        'worker_storage': ['server-storage-03'], 'worker': ['server-worker-04'],
        'infra': ['server-infra-05'],
    }
    assert len(result['node_endpoints']) == 5
    assert result['node_endpoints']['server-gpu-02']['bmc_fqdn'] == 'gpu-02.management.example.com'
    assert result['node_endpoints']['server-worker-04']['rhcos_fqdn'] == 'worker-04.cluster-01.example.com'
    # Matching generated names are retained on a subsequent explicitly forced regeneration.
    monkeypatch.setattr(sys, 'argv', ['init-ocp-settings.py', '--arch', 'demo', '--site', 'site', '--force'])
    INIT.main()
    assert yaml.safe_load(output.read_text())['node_endpoints'] == result['node_endpoints']


def test_partial_roles_rejected_and_blank_roles_retain_legacy(tmp_path):
    path = tmp_path / 'workbook.xlsx'
    assignments = [('support', 'server-01', 'control', 'Yes', 'Yes'),
                   ('gpu', 'server-02', None, 'Yes', 'No')]
    groups = {'support': ['server-01'], 'nodes': ['server-02']}
    _role_workbook(path, assignments)
    with pytest.raises(ValueError, match='exactly every enabled'):
        INIT.read_workbook_roles(path, groups)
    _role_workbook(path, [(f, n, None, e, m) for f, n, _, e, m in assignments])
    roles, subdomain = INIT.read_workbook_roles(path, groups)
    assert roles is None and subdomain == 'management'
    assert INIT.build_node_roles(groups) == {'worker_gpu': ['server-02'], 'control_plane': ['server-01']}


def test_role_generation_preserves_file_on_mapping_conflict(tmp_path):
    output = tmp_path / 'ocp-settings.yml'
    original = {'node_endpoints': {'server-01': {
        'bmc_fqdn': 'custom.mgmt-01.example.com', 'rhcos_fqdn': 'custom.cluster-01.example.com'}}}
    output.write_text(yaml.safe_dump(original))
    before = output.read_bytes()
    with pytest.raises(ValueError, match='conflict'):
        INIT.write_settings(output, 'demo', 'site', '  control_plane:\n    - server-01',
                            ocp={'ocp_cluster_name': 'cluster-01', 'ocp_cluster_domain': 'example.com'},
                            workbook_roles={'server-01': 'control'})
    assert output.read_bytes() == before


def test_validator_rejects_partial_roles_and_generated_collisions(tmp_path):
    from validate_excel import validate_nodes, ValidationResult
    import openpyxl
    path = tmp_path / 'workbook.xlsx'
    _role_workbook(path, [('support', 'server-a-01', 'worker', 'Yes', 'Yes'),
                         ('support', 'server-b-01', 'worker', 'Yes', 'Yes')])
    wb = openpyxl.load_workbook(path)
    # This small fixture tests naming validation only, supplying its required IP header.
    wb['Nodes']['F1'] = 'Mgmt IP Address'
    settings = {'ocp_cluster_name': 'cluster-01', 'ocp_cluster_domain': 'example.com'}
    result = ValidationResult()
    validate_nodes(wb['Nodes'], result, settings=settings)
    assert any('Duplicate generated endpoint' in error for error in result.errors)
    wb['Nodes']['E3'] = None
    result = ValidationResult()
    validate_nodes(wb['Nodes'], result, settings=settings)
    assert any('Role must cover exactly' in error for error in result.errors)
    wb.close()
